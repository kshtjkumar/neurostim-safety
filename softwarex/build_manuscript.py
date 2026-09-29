"""Build the SoftwareX manuscript draft as an editable .docx and a readable .pdf.

Source: manuscript_draft.md (this folder). Figures: figures/*.png (run
figures/make_figures.py first). The .docx is built on Elsevier's official template
(softwarex-osp-template.docx, v6, March 2026 -- download it next to this script from
https://legacyfileshare.elsevier.com/promis_misc/softwarex-osp-template.docx); its
instruction text is removed and its styles are kept. [AUTHOR: ...] markers are
highlighted in both outputs.

Requires python-docx and reportlab:  python build_manuscript.py
"""

from __future__ import annotations

import re
from pathlib import Path

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "manuscript_draft.md"
TEMPLATE = HERE / "softwarex-osp-template.docx"
FIGURES = {
    "1": HERE / "figures" / "figure1_architecture.png",
    "2": HERE / "figures" / "figure2_dbs3389_summary.png",
}
AUTHOR = re.compile(r"(\[AUTHOR:.*?\])", re.S)
FIGURE = re.compile(r"^\[Figure (\d) here\. Caption: (.*)\]$", re.S)


def blocks(text: str) -> list[tuple[str, object]]:
    """Split the Markdown into (kind, payload) blocks: h1/h2/h3, p, ul, table, code, fig."""
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S).strip()
    out: list[tuple[str, object]] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            out.append(("code", "\n".join(lines[i + 1 : j])))
            i = j + 1
            continue
        m = re.match(r"^(#{1,3}) (.*)", line)
        if m:
            out.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            out.append(("table", rows))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and (lines[i].startswith("- ") or lines[i].startswith("  ")):
                if lines[i].startswith("- "):
                    items.append(lines[i][2:].strip())
                else:
                    items[-1] += " " + lines[i].strip()
                i += 1
            out.append(("ul", items))
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||- |```)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        joined = " ".join(para)
        fm = FIGURE.match(joined)
        out.append(("fig", (fm.group(1), fm.group(2))) if fm else ("p", joined))
    return out


def inline_runs(text: str) -> list[tuple[str, str]]:
    """(style, text) runs: 'author' for [AUTHOR ...], 'code', 'bold', 'italic', 'plain'."""
    runs: list[tuple[str, str]] = []
    for part in AUTHOR.split(text):
        if AUTHOR.fullmatch(part):
            runs.append(("author", part))
            continue
        for tok in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)", part):
            if not tok:
                continue
            if tok.startswith("`"):
                runs.append(("code", tok[1:-1]))
            elif tok.startswith("**"):
                runs.append(("bold", tok[2:-2]))
            elif tok.startswith("*"):
                runs.append(("italic", tok[1:-1]))
            else:
                runs.append(("plain", tok))
    return runs


def build_docx(parts, target: Path) -> None:
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Mm, Pt

    from copy import deepcopy

    doc = docx.Document(str(TEMPLATE))
    body = doc.element.body
    # Elsevier's own code-metadata table (C1-C8), kept with its formatting and filled below.
    metadata_tbl = deepcopy(doc.tables[0]._tbl)
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)

    def fill(par, text):
        for style, chunk in inline_runs(text):
            run = par.add_run(chunk)
            if style == "author":
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif style == "code":
                run.font.name = "Courier New"
                run.font.size = Pt(9)
            elif style == "bold":
                run.bold = True
            elif style == "italic":
                run.italic = True

    for kind, payload in parts:
        if kind in ("h1", "h2", "h3"):
            level = {"h1": 0, "h2": 1, "h3": 2}[kind]
            if level == 0:
                par = doc.add_paragraph(style="Title")
                fill(par, payload)
            else:
                par = doc.add_heading(level=level)
                fill(par, payload)
        elif kind == "p":
            fill(doc.add_paragraph(style="Normal"), payload)
        elif kind == "ul":
            for item in payload:
                fill(doc.add_paragraph(style="List Paragraph"), "• " + item)
        elif kind == "code":
            par = doc.add_paragraph(style="Normal")
            run = par.add_run(payload)
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
        elif kind == "table":
            from docx.table import Table as DocxTable

            body.insert(len(body) - 1, metadata_tbl)  # before the final sectPr
            table = DocxTable(metadata_tbl, doc._body)
            for r, cells in enumerate(payload[1:], start=1):  # skip our header row
                cell = table.cell(r, len(table.columns) - 1)
                for par in cell.paragraphs[1:]:
                    par._element.getparent().remove(par._element)
                cell.paragraphs[0].text = ""
                fill(cell.paragraphs[0], cells[-1])
        elif kind == "fig":
            number, caption = payload
            doc.add_picture(str(FIGURES[number]), width=Mm(165))
            fill(doc.add_paragraph(style="Normal"), f"**Fig. {number}.** {caption}")
    doc.save(str(target))


def build_pdf(parts, target: Path) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, Paragraph, Preformatted, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.5,
                          leading=13, spaceAfter=5)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10)
    styles = {
        "h1": ParagraphStyle("t", parent=ss["Title"], fontSize=15, leading=19),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, spaceBefore=10),
        "h3": ParagraphStyle("h3", parent=ss["Heading3"], fontSize=10.5, spaceBefore=6),
    }

    def markup(text: str) -> str:
        out = []
        for style, chunk in inline_runs(text):
            chunk = escape(chunk)
            if style == "author":
                out.append(f'<font backColor="#FFF3A0">{chunk}</font>')
            elif style == "code":
                out.append(f'<font face="Courier">{chunk}</font>')
            elif style == "bold":
                out.append(f"<b>{chunk}</b>")
            elif style == "italic":
                out.append(f"<i>{chunk}</i>")
            else:
                out.append(chunk)
        return "".join(out)

    story = []
    for kind, payload in parts:
        if kind in styles:
            story.append(Paragraph(markup(payload), styles[kind]))
        elif kind == "p":
            story.append(Paragraph(markup(payload), body))
        elif kind == "ul":
            for item in payload:
                story.append(Paragraph(markup(item), body, bulletText="•"))
        elif kind == "code":
            story.append(Preformatted(payload, ParagraphStyle("code", fontName="Courier",
                                                               fontSize=8, leading=10)))
        elif kind == "table":
            data = [[Paragraph(markup(c), small) for c in row] for row in payload]
            table = Table(data, colWidths=[12 * mm, 52 * mm, 106 * mm], repeatRows=1)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [table, Spacer(1, 6)]
        elif kind == "fig":
            number, caption = payload
            path = FIGURES[number]
            from PIL import Image as PILImage

            w, h = PILImage.open(path).size
            width = 170 * mm
            story += [Image(str(path), width=width, height=width * h / w),
                      Paragraph(markup(f"**Fig. {number}.** {caption}"), small), Spacer(1, 6)]
    SimpleDocTemplate(str(target), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                      topMargin=18 * mm, bottomMargin=18 * mm,
                      title="neurostim-safety (SoftwareX draft)").build(story)


if __name__ == "__main__":
    parts = blocks(SOURCE.read_text(encoding="utf-8"))
    build_docx(parts, HERE / "manuscript_draft.docx")
    build_pdf(parts, HERE / "manuscript_draft.pdf")
    print("written: manuscript_draft.docx, manuscript_draft.pdf")
